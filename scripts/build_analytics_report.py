#!/usr/bin/env python3
"""Xuất báo cáo Analytics + Điểm mù ra Word (.docx) và Excel (.xlsx).

Nguồn: `outputs/reports/blindspots.json` (đã phân tích từ analytics thật) +
`vendor/youtube-analytics-dashboard/analytics_latest.json` (private metrics OAuth) +
`vendor/youtube-analytics-dashboard/history.csv` (public stats).

Tách riêng khỏi `build_docx_report.py` (báo cáo đối thủ) vì nội dung khác hẳn:
đây là số liệu kênh mình, không phải phân tích đối thủ.

Usage:
    python scripts/build_analytics_report.py
"""
from __future__ import annotations

import csv
import json
import re
import sys
from datetime import date
from html import unescape
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from console_utf8 import ensure_utf8_console  # noqa: E402

ensure_utf8_console()

ROOT = Path(__file__).resolve().parents[1]
VENDOR = ROOT / "vendor" / "youtube-analytics-dashboard"
OUT = ROOT / "outputs" / "reports"
OWNER_ID = "UCBZ7LaffmEPv91sWcfroJdQ"

SEV_COLOR = {
    "CRITICAL": RGBColor(0xB9, 0x1C, 0x1C),
    "HIGH": RGBColor(0xC2, 0x54, 0x0C),
    "MEDIUM": RGBColor(0x9A, 0x6A, 0x00),
    "INFO": RGBColor(0x1F, 0x4E, 0x79),
}
SEV_FILL = {"CRITICAL": "F8D7DA", "HIGH": "FDE2CC",
            "MEDIUM": "FFF3CD", "INFO": "D6E4F0"}


def load_json(p: Path):
    if not p.exists():
        return None
    return json.loads(p.read_text(encoding="utf-8"))


def load_csv(p: Path) -> list[dict]:
    if not p.exists():
        return []
    with p.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def to_int(v, d=0) -> int:
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return d


def shade(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(__import__("docx").oxml.ns.qn("w:fill"), hex_color)
    tc_pr.append(shd)


def build_docx(blind: dict, an: dict, hist: list[dict]) -> Path:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(4)
    for lvl in range(1, 5):
        try:
            doc.styles[f"Heading {lvl}"].font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
        except KeyError:
            pass

    h = doc.add_heading("BÁO CÁO ANALYTICS & ĐIỂM MÙ KÊNH", level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Kênh: @azzammastertradinggold  •  Ngách: XAUUSD / Forex").bold = True
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    w = blind.get("window", {})
    p2.add_run(f"Cửa sổ phân tích: {w.get('start')} → {w.get('end')} "
               f"({w.get('days')} ngày)\nSinh: {blind.get('generated_at','')}")
    doc.add_paragraph()

    # --- Cảnh báo quan trọng nhất ---
    bot = blind.get("bot_traffic", {})
    if to_int(bot.get("views")) > 0:
        doc.add_heading("CẢNH BÁO — Traffic nghi bot", level=1)
        warn = doc.add_paragraph()
        r = warn.add_run(
            f"{to_int(bot.get('views')):,} views ({bot.get('pct_of_channel')}% tổng view) "
            f"đến từ referrer không phải site thật. Đây là dấu hiệu view mua/view farm. "
            f"YouTube có thể xoá view hoặc phạt kênh theo Fake Engagement Policy. "
            f"Cần điều tra ngay nguồn gốc traffic này.")
        r.bold = True
        r.font.color.rgb = SEV_COLOR["CRITICAL"]
        doc.add_paragraph()

    # --- Số liệu nền ---
    doc.add_heading("1. Số liệu nền (28 ngày)", level=1)
    m = blind.get("metrics", {})
    dv = blind.get("derived", {})
    latest = None
    for r in hist:
        if r.get("channel_id") == OWNER_ID:
            if latest is None or r.get("date", "") > latest.get("date", ""):
                latest = r

    rows = [
        ("Views", f"{to_int(m.get('views')):,}", "trong cửa sổ"),
        ("Watch time", f"{dv.get('watch_hours')} giờ", "tổng thời lượng xem"),
        ("View trung bình", f"{to_int(m.get('averageViewDuration'))}s",
         f"{m.get('averageViewPercentage')}% độ dài video"),
        ("Subscribers", f"+{to_int(m.get('subscribersGained'))} / "
                        f"-{to_int(m.get('subscribersLost'))}",
         f"ròng {dv.get('sub_net'):+d}"),
        ("Sub conversion", f"{dv.get('sub_conversion_pct')}%", "tham chiếu ngành 0.5–2%"),
        ("Likes", f"{to_int(m.get('likes')):,}", f"like rate {dv.get('like_rate_pct')}%"),
        ("Comments", f"{to_int(m.get('comments'))}", f"comment rate {dv.get('comment_rate_pct')}%"),
        ("Shares", f"{to_int(m.get('shares'))}", f"share rate {dv.get('share_rate_pct')}%"),
    ]
    if latest:
        rows += [
            ("Tổng sub kênh", f"{to_int(latest.get('subs')):,}", "tích luỹ"),
            ("Tổng view kênh", f"{to_int(latest.get('views')):,}", "tích luỹ"),
            ("Số video", f"{to_int(latest.get('videos')):,}", "tích luỹ"),
        ]
    t = doc.add_table(rows=1, cols=3)
    t.style = "Light Grid Accent 1"
    hdr = t.rows[0].cells
    for i, txt in enumerate(("Chỉ số", "Giá trị", "Ghi chú")):
        hdr[i].text = txt
        for par in hdr[i].paragraphs:
            for run in par.runs:
                run.bold = True
    for a, b, c in rows:
        cells = t.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = a, b, c
    doc.add_paragraph()

    # --- Nguồn traffic ---
    ts = an.get("traffic_sources", []) if an else []
    if ts:
        doc.add_heading("2. Nguồn traffic", level=1)
        tot = sum(to_int(r[1]) for r in ts) or 1
        t2 = doc.add_table(rows=1, cols=4)
        t2.style = "Light Grid Accent 1"
        for i, txt in enumerate(("Nguồn", "Views", "Tỷ lệ", "Watch (phút)")):
            t2.rows[0].cells[i].text = txt
            for par in t2.rows[0].cells[i].paragraphs:
                for run in par.runs:
                    run.bold = True
        for r in ts:
            v = to_int(r[1])
            c = t2.add_row().cells
            c[0].text = str(r[0])
            c[1].text = f"{v:,}"
            c[2].text = f"{v/tot*100:.1f}%"
            c[3].text = f"{to_int(r[2]):,}"
        doc.add_paragraph()

    # --- Nguồn ngoài ---
    ex = an.get("external_urls", []) if an else []
    if ex:
        doc.add_heading("3. Nguồn ngoài (EXT_URL) — kiểm tra bot", level=1)
        BOT = ("seofast", "playbot", "viewbot", "like4like", "sub4sub")
        t3 = doc.add_table(rows=1, cols=3)
        t3.style = "Light Grid Accent 1"
        for i, txt in enumerate(("Referrer", "Views", "Cảnh báo")):
            t3.rows[0].cells[i].text = txt
            for par in t3.rows[0].cells[i].paragraphs:
                for run in par.runs:
                    run.bold = True
        for r in ex:
            is_bot = any(b in str(r[0]).lower() for b in BOT)
            c = t3.add_row().cells
            c[0].text = str(r[0])
            c[1].text = f"{to_int(r[1]):,}"
            c[2].text = "NGHI BOT" if is_bot else ""
            if is_bot:
                for par in c[2].paragraphs:
                    for run in par.runs:
                        run.bold = True
                        run.font.color.rgb = SEV_COLOR["CRITICAL"]
        doc.add_paragraph()

    # --- Điểm mù ---
    fs = blind.get("findings", [])
    doc.add_heading(f"4. Điểm mù phát hiện được ({len(fs)})", level=1)
    sev_order = {"CRITICAL": 0, "HIGH": 1, "MEDIUM": 2, "INFO": 3}
    for i, f in enumerate(sorted(fs, key=lambda x: sev_order.get(x.get("severity"), 9)), 1):
        sev = f.get("severity", "")
        hh = doc.add_heading(f"{i}. [{sev}] {unescape(f.get('title',''))}", level=2)
        for run in hh.runs:
            run.font.color.rgb = SEV_COLOR.get(sev, RGBColor(0, 0, 0))
        par = doc.add_paragraph()
        par.add_run(unescape(f.get("detail", "")))
        ev = f.get("evidence")
        if ev:
            ep = doc.add_paragraph()
            ep.add_run("Bằng chứng: ").italic = True
            ep.add_run(json.dumps(ev, ensure_ascii=False)).font.name = "Consolas"
        doc.add_paragraph()

    doc.add_page_break()
    doc.add_heading("5. Phương pháp & giới hạn", level=1)
    for line in [
        "Nguồn dữ liệu: YouTube Analytics API v2 (OAuth 2.0, scope yt-analytics.readonly) "
        "cho private metrics; YouTube Data API v3 cho public stats.",
        f"Channel id tường minh: {OWNER_ID} — KHÔNG dùng channel==MINE vì kênh là Brand "
        f"Account, MINE sẽ trả số liệu của kênh cá nhân (sai hoàn toàn).",
        "Mọi con số đều tính từ file JSON/CSV thật trong repo. Không có số nhập tay.",
        "Cửa sổ analytics có độ trễ ~2 ngày (Google không trả dữ liệu 2 ngày gần nhất).",
        "Phát hiện bot dựa trên danh sách marker referrer — có thể bỏ sót nguồn mới. "
        "Đây là tín hiệu điều tra, không phải kết luận vi phạm.",
        "Không có cam kết lợi nhuận, không có kết quả trading mô phỏng trong báo cáo này.",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    path = OUT / "AZZAM_ANALYTICS_DIEM_MU.docx"
    doc.save(path)
    return path


def build_xlsx(blind: dict, an: dict, hist: list[dict]) -> Path:
    wb = Workbook()
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F3864")
    head_font = Font(bold=True, color="FFFFFF", size=10)

    def style_header(ws, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border
        ws.freeze_panes = "A2"

    def autosize(ws, widths):
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # --- Sheet 1: Tổng quan ---
    ws = wb.active
    ws.title = "01_TONG_QUAN"
    m, dv = blind.get("metrics", {}), blind.get("derived", {})
    w = blind.get("window", {})
    ws.append(["Chỉ số", "Giá trị", "Đơn vị", "Ghi chú"])
    data = [
        ("Cửa sổ bắt đầu", w.get("start"), "ngày", "analytics trễ ~2 ngày"),
        ("Cửa sổ kết thúc", w.get("end"), "ngày", ""),
        ("Số ngày", w.get("days"), "ngày", ""),
        ("Views", to_int(m.get("views")), "views", ""),
        ("Watch time", dv.get("watch_hours"), "giờ", ""),
        ("View trung bình", to_int(m.get("averageViewDuration")), "giây", ""),
        ("% video xem được", m.get("averageViewPercentage"), "%", ""),
        ("Subs gained", to_int(m.get("subscribersGained")), "sub", ""),
        ("Subs lost", to_int(m.get("subscribersLost")), "sub", ""),
        ("Sub ròng", dv.get("sub_net"), "sub", ""),
        ("Sub conversion", dv.get("sub_conversion_pct"), "%", "ngành 0.5–2%"),
        ("Likes", to_int(m.get("likes")), "like", ""),
        ("Like rate", dv.get("like_rate_pct"), "%", ""),
        ("Comments", to_int(m.get("comments")), "comment", ""),
        ("Comment rate", dv.get("comment_rate_pct"), "%", ""),
        ("Shares", to_int(m.get("shares")), "share", ""),
        ("Share rate", dv.get("share_rate_pct"), "%", ""),
        ("Traffic nghi bot", to_int(blind.get("bot_traffic", {}).get("views")), "views", ""),
        ("% bot trên tổng view", blind.get("bot_traffic", {}).get("pct_of_channel"), "%", ""),
    ]
    for r in data:
        ws.append(list(r))
    style_header(ws, 4)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = border
    autosize(ws, [26, 16, 12, 26])

    # --- Sheet 2: Điểm mù ---
    ws2 = wb.create_sheet("02_DIEM_MU")
    ws2.append(["#", "Mức độ", "Tiêu đề", "Chi tiết", "Bằng chứng (JSON)"])
    sev_order = {"CRITICAL": 0, "HIGH": 1, "MEDIUM": 2, "INFO": 3}
    for i, f in enumerate(sorted(blind.get("findings", []),
                                 key=lambda x: sev_order.get(x.get("severity"), 9)), 1):
        ws2.append([i, f.get("severity"), f.get("title"), f.get("detail"),
                    json.dumps(f.get("evidence"), ensure_ascii=False)])
    style_header(ws2, 5)
    for row in ws2.iter_rows(min_row=2):
        sev = row[1].value
        fill = SEV_FILL.get(sev)
        for c in row:
            c.border = border
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if fill:
                c.fill = PatternFill("solid", fgColor=fill)
        row[1].font = Font(bold=True)
    autosize(ws2, [5, 11, 46, 78, 60])

    # --- Sheet 3: Traffic ---
    ws3 = wb.create_sheet("03_TRAFFIC")
    ws3.append(["Nguồn", "Views", "Tỷ lệ %", "Watch (phút)", "Giây/view"])
    ts = an.get("traffic_sources", []) if an else []
    tot = sum(to_int(r[1]) for r in ts) or 1
    for r in ts:
        v, mins = to_int(r[1]), to_int(r[2])
        ws3.append([r[0], v, round(v / tot * 100, 2), mins,
                    round(mins * 60 / v, 1) if v else 0])
    style_header(ws3, 5)
    for row in ws3.iter_rows(min_row=2):
        for c in row:
            c.border = border
        row[1].number_format = "#,##0"
        row[3].number_format = "#,##0"
    autosize(ws3, [22, 12, 10, 14, 11])

    # --- Sheet 4: Nguồn ngoài (bot) ---
    ws4 = wb.create_sheet("04_NGUON_NGOAI")
    ws4.append(["Referrer", "Views", "Watch (phút)", "Nghi bot"])
    BOT = ("seofast", "playbot", "viewbot", "like4like", "sub4sub")
    for r in (an.get("external_urls", []) if an else []):
        is_bot = any(b in str(r[0]).lower() for b in BOT)
        ws4.append([r[0], to_int(r[1]), to_int(r[2]), "CÓ" if is_bot else ""])
    style_header(ws4, 4)
    for row in ws4.iter_rows(min_row=2):
        for c in row:
            c.border = border
        row[1].number_format = "#,##0"
        if row[3].value == "CÓ":
            row[3].font = Font(bold=True, color="B91C1C")
            for c in row:
                c.fill = PatternFill("solid", fgColor="F8D7DA")
    autosize(ws4, [40, 12, 14, 11])

    # --- Sheet 5: Daily ---
    ws5 = wb.create_sheet("05_THEO_NGAY")
    ws5.append(["Ngày", "Views", "Watch (phút)", "% tổng view"])
    daily = an.get("daily", []) if an else []
    tot_v = sum(to_int(r[1]) for r in daily) or 1
    for r in daily:
        ws5.append([r[0], to_int(r[1]), round(float(r[2] or 0), 1),
                    round(to_int(r[1]) / tot_v * 100, 2)])
    style_header(ws5, 4)
    for row in ws5.iter_rows(min_row=2):
        for c in row:
            c.border = border
        row[1].number_format = "#,##0"
        row[2].number_format = "#,##0.0"
    autosize(ws5, [14, 12, 14, 12])

    # --- Sheet 6: Địa lý ---
    ws6 = wb.create_sheet("06_DIA_LY")
    ws6.append(["Quốc gia", "Views", "Watch (phút)", "Tỷ lệ %", "Giây/view"])
    geo = an.get("geography", []) if an else []
    gt = sum(to_int(r[1]) for r in geo) or 1
    for r in geo:
        v, mins = to_int(r[1]), to_int(r[2])
        ws6.append([r[0], v, mins, round(v / gt * 100, 2),
                    round(mins * 60 / v, 0) if v else 0])
    style_header(ws6, 5)
    for row in ws6.iter_rows(min_row=2):
        for c in row:
            c.border = border
        row[1].number_format = "#,##0"
        row[2].number_format = "#,##0"
    autosize(ws6, [12, 12, 14, 10, 11])

    # --- Sheet 7: Thiết bị ---
    ws7 = wb.create_sheet("07_THIET_BI")
    ws7.append(["Thiết bị", "Views", "Tỷ lệ %"])
    dev = an.get("devices", []) if an else []
    dt = sum(to_int(r[1]) for r in dev) or 1
    for r in dev:
        ws7.append([r[0], to_int(r[1]), round(to_int(r[1]) / dt * 100, 2)])
    style_header(ws7, 3)
    for row in ws7.iter_rows(min_row=2):
        for c in row:
            c.border = border
        row[1].number_format = "#,##0"
    autosize(ws7, [14, 12, 10])

    # --- Sheet 8: Top video ---
    ws8 = wb.create_sheet("08_TOP_VIDEO")
    ws8.append(["Video ID", "Watch (phút)", "Views", "% xem được", "Subs gained", "Tiêu đề"])
    for r in (an.get("top_videos", []) if an else []):
        ws8.append([r[0], to_int(r[1]), to_int(r[2]), r[3],
                    to_int(r[4]) if len(r) > 4 else 0, r[5] if len(r) > 5 else ""])
    style_header(ws8, 6)
    for row in ws8.iter_rows(min_row=2):
        for c in row:
            c.border = border
            c.alignment = Alignment(vertical="top", wrap_text=(c.column == 6))
        row[1].number_format = "#,##0"
        row[2].number_format = "#,##0"
    autosize(ws8, [14, 14, 10, 12, 12, 62])

    # --- Sheet 9: Search terms ---
    ws9 = wb.create_sheet("09_SEARCH_TERMS")
    ws9.append(["Từ khoá tìm kiếm", "Views"])
    for r in (an.get("search_terms", []) if an else []):
        ws9.append([r[0], to_int(r[1])])
    style_header(ws9, 2)
    for row in ws9.iter_rows(min_row=2):
        for c in row:
            c.border = border
        row[1].number_format = "#,##0"
    autosize(ws9, [40, 10])

    # --- Sheet 10: Lịch sử kênh ---
    ws10 = wb.create_sheet("10_LICH_SU_KENH")
    ws10.append(["Ngày", "Kênh", "Subs", "Tổng views", "Số video"])
    for r in sorted(hist, key=lambda x: (x.get("date", ""), x.get("title", "")),
                    reverse=True):
        ws10.append([r.get("date"), r.get("title"), to_int(r.get("subs")),
                     to_int(r.get("views")), to_int(r.get("videos"))])
    style_header(ws10, 5)
    for row in ws10.iter_rows(min_row=2):
        for c in row:
            c.border = border
        row[2].number_format = "#,##0"
        row[3].number_format = "#,##0"
        row[4].number_format = "#,##0"
    autosize(ws10, [12, 24, 10, 13, 11])

    path = OUT / "AZZAM_ANALYTICS_DIEM_MU.xlsx"
    wb.save(path)
    return path


def main() -> int:
    blind = load_json(OUT / "blindspots.json")
    if not blind:
        print("Thiếu outputs/reports/blindspots.json — chạy scripts/analyze_blindspots.py trước.",
              file=sys.stderr)
        return 1
    an = load_json(VENDOR / "analytics_latest.json")
    hist = load_csv(VENDOR / "history.csv")

    d = build_docx(blind, an or {}, hist)
    x = build_xlsx(blind, an or {}, hist)

    print(f"Saved: {d}  ({d.stat().st_size:,} bytes)")
    print(f"Saved: {x}  ({x.stat().st_size:,} bytes)")

    # Verify bằng cách đọc lại — không tin việc save thành công.
    doc = Document(str(d))
    print(f"\n  DOCX: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables")
    wb = Workbook()
    wb = __import__("openpyxl").load_workbook(str(x), data_only=True)
    print(f"  XLSX: {len(wb.sheetnames)} sheets")
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"    {name}: {ws.max_row - 1} data rows")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
