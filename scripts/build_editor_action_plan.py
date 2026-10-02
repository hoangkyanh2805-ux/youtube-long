"""
Action plan cho đội edit + bộ comment tương tác (seeding an toàn chính sách).

Writes:
  outputs/reports/AZZAM_ACTION_PLAN_EDITOR.docx
  outputs/reports/AZZAM_ACTION_PLAN_EDITOR.xlsx

Nội dung:
  - KPI ngày: 3 Short + 1 Long 8 phút
  - Phân công + workflow sản xuất
  - Nguồn dữ liệu pain point (map từng KPI về file + cột cụ thể)
  - Bộ comment tương tác: gợi ý tranh cãi / gây tò mò / hỏi ngược
  - Ranh giới chính sách YouTube: được gì, không được gì

QUAN TRỌNG — chính sách YouTube:
  Seeding bằng tài khoản giả, mua engagement, hoặc spam link là vi phạm
  (Fake Engagement Policy + Spam/Deceptive Practices). Bộ comment trong file
  này được thiết kế để CHÍNH KÊNH đăng công khai (pinned comment / Community
  tab) và để KHÁN GIẢ THẬT dùng — không phải để tài khoản ảo đi rải.
"""
from __future__ import annotations

import csv
import json
from datetime import date
from html import unescape
from pathlib import Path
import sys

# Force UTF-8 console: without this, printing Vietnamese crashes with
# UnicodeEncodeError when stdout is a pipe (dagu/cron/CI).
sys.path.insert(0, str(Path(__file__).resolve().parent))
from console_utf8 import ensure_utf8_console  # noqa: E402
ensure_utf8_console()

OUT = Path("outputs/reports")
OUT.mkdir(parents=True, exist_ok=True)

KPI = {"shorts_per_day": 3, "long_per_day": 1, "long_minutes": 8}


# ── Data ────────────────────────────────────────────────────────────────────

def load_csv(path: str) -> list[dict]:
    p = Path(path)
    if not p.exists():
        return []
    with p.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def load_json(path: str):
    p = Path(path)
    if not p.exists():
        return None
    return json.loads(p.read_text(encoding="utf-8"))


def dedupe(rows: list[dict]) -> list[dict]:
    out, seen = [], set()
    for r in rows:
        cid = r.get("comment_id", "")
        key = cid or f"{r.get('content_url','')}|{r.get('comment_text','')[:80]}"
        if key in seen:
            continue
        seen.add(key)
        out.append(r)
    return out


def to_float(v) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


# ── KPI & workflow ──────────────────────────────────────────────────────────

def build_kpi_rows(pain_all: list[dict], backlog: list[dict]) -> list[list]:
    return [
        ["SHORT 1", "Hook pain point (số 1)", "00:15-00:35", "1 pain point duy nhất",
         "3s hook >60% giữ chân", "06_KEYWORD_CHINH + 09_PAINPOINT_ALL"],
        ["SHORT 2", "Hook pain point (số 2)", "00:15-00:35", "1 pain point duy nhất",
         "3s hook >60% giữ chân", "09_PAINPOINT_ALL (score ≥ 7)"],
        ["SHORT 3", "Hook pain point (số 3)", "00:15-00:35", "1 pain point duy nhất",
         "3s hook >60% giữ chân", "09_PAINPOINT_ALL (score ≥ 6)"],
        ["LONG 1", "Mini-course 8 phút", "08:00 (7:30-8:30)", "1 pillar duy nhất",
         "Avg view duration >40%", "04_SALES_ANGLES + 05_CONTENT_BACKLOG"],
    ]


def build_workflow_rows() -> list[list]:
    return [
        ["1. Chọn nguyên liệu", "Editor", "Mở 09_PAINPOINT_ALL, lọc score cao, "
         "chọn 3 comment khác nhóm pain", "3 pain point đã chốt + comment_id"],
        ["2. Viết hook 3 giây", "Editor", "Câu đầu phải nêu thẳng nỗi đau, không "
         "intro, không logo", "Script 3 short (mỗi cái ≤35s)"],
        ["3. Quay/dựng Short", "Editor", "Dọc 9:16, phụ đề burn-in, cắt mọi khoảng lặng",
         "3 file .mp4 dọc"],
        ["4. Dựng Long 8 phút", "Editor", "Cấu trúc: vấn đề → nguyên nhân → "
         "giải pháp → checklist → CTA", "1 file .mp4 ngang"],
        ["5. Packaging", "Editor", "Title theo 08_TITLE_PATTERN, thumbnail 3 yếu tố "
         "(mặt/số/chữ ngắn)", "Title + thumbnail mỗi video"],
        ["6. SEO", "Editor", "Description có timestamp + tag theo 06/07_KEYWORD",
         "Description + tag đã điền"],
        ["7. Kiểm tra guardrail", "Editor + duyệt", "Không cam kết lợi nhuận, không "
         "trade giả, affiliate có disclosure", "Checklist ký duyệt"],
        ["8. Đăng + pinned comment", "Chủ kênh", "Đăng theo 11_LICH_DANG, ghim "
         "comment tương tác từ sheet 12", "Link video + comment đã ghim"],
    ]


def build_source_rows() -> list[list]:
    """Map từng loại dữ liệu về file + cột cụ thể để đội edit tự lấy."""
    return [
        ["Pain point có score", "09_PAINPOINT_ALL", "Score, Categories, Nội dung comment, "
         "URL nguồn, Comment ID", "Chọn hook Short + chủ đề Long"],
        ["Comment gốc chưa lọc", "10_COMMENTS_RAW", "Nội dung comment, Likes, URL",
         "Tìm thêm góc ngoài danh sách đã lọc"],
        ["Góc bán hàng theo nhóm pain", "04_SALES_ANGLES", "Pain cluster, Sales angle, "
         "Big promise, Proof hook, CTA", "Viết script Long theo pillar"],
        ["Hàng đợi content", "05_CONTENT_BACKLOG", "Ưu tiên, Format, Tiêu đề nháp, CTA",
         "Biết làm video nào trước"],
        ["Từ khoá chính", "06_KEYWORD_CHINH", "Từ khoá, Tần suất, Ưu tiên",
         "Đặt title + tag"],
        ["Từ khoá dài", "07_KEYWORD_DAI", "Cụm từ, Tần suất", "Viết hook theo ngôn ngữ khán giả"],
        ["Pattern tiêu đề đối thủ", "08_TITLE_PATTERN", "Pattern, Số video dùng",
         "Công thức đặt title"],
        ["Video đối thủ thắng", "03_VIDEO_INVENTORY", "Tiêu đề, Views, Nhóm thời lượng, "
         "Like rate %", "Tham chiếu format + packaging"],
    ]


# ── Comment tương tác (an toàn chính sách) ──────────────────────────────────

SEED_COMMENTS = [
    # (nhóm, loại, câu comment, dùng ở đâu, rủi ro chính sách)
    ("Risk management", "Gợi ý tranh cãi nhẹ",
     "Mình biết sẽ có người phản đối: risk 1%/lệnh là quá thấp. Nhưng bạn thử tính "
     "xem bao nhiêu lệnh thua liên tiếp thì tài khoản về 0 với 5%/lệnh?",
     "Pinned comment video về risk", "Thấp — không link, không hứa lợi nhuận"),
    ("Risk management", "Hỏi ngược",
     "Câu hỏi thật: nếu chỉ được chọn 1 thứ để sửa trong cách trade hiện tại, bạn "
     "sửa entry hay sửa risk? Mình chọn risk, nhưng tò mò ý mọi người.",
     "Pinned comment", "Thấp"),
    ("Risk management", "Gây tò mò",
     "Có một con số trong video này mà 90% người mới tính sai. Bạn thử đoán xem "
     "là số nào trước khi xem hết?",
     "Pinned comment", "Thấp — không clickbait sai sự thật"),

    ("Strategy / overcomplicating", "Gợi ý tranh cãi nhẹ",
     "Nói thẳng: nếu chart của bạn có hơn 3 indicator, bạn đang tự làm khó mình. "
     "Ai không đồng ý thì nói lý do, mình đọc hết.",
     "Pinned comment video về đơn giản hoá", "Thấp — quan điểm, không xúc phạm"),
    ("Strategy / overcomplicating", "Hỏi ngược",
     "Bạn đang dùng bao nhiêu indicator? Trả lời 1 con số thôi. Mình cá là phần "
     "lớn sẽ nói 5-10, và đó là vấn đề.",
     "Pinned comment", "Thấp"),
    ("Strategy / overcomplicating", "Gây tò mò",
     "Mình đã xoá 9/10 indicator của mình và kết quả đổi theo hướng không ngờ. "
     "Bạn nghĩ xoá indicator sẽ làm tệ hơn hay tốt hơn?",
     "Pinned comment", "Thấp"),

    ("Entry timing", "Gợi ý tranh cãi nhẹ",
     "Vào lệnh sớm 3 nến nghe nhỏ, nhưng đó là khác biệt giữa lệnh thắng và lệnh "
     "bị quét stop loss. Bạn thuộc team vào sớm hay team chờ xác nhận?",
     "Pinned comment video entry", "Thấp"),
    ("Entry timing", "Hỏi ngược",
     "Bạn có đang chờ nến đóng cửa trước khi vào lệnh không? Nếu không, lý do là gì?",
     "Pinned comment", "Thấp"),

    ("Psychology / discipline", "Gợi ý tranh cãi nhẹ",
     "Ý kiến gây tranh cãi: revenge trading không phải vấn đề cảm xúc, mà là vấn đề "
     "quy trình. Bạn nghĩ sao?",
     "Pinned comment video tâm lý", "Thấp"),
    ("Psychology / discipline", "Hỏi ngược",
     "Bao nhiêu lần bạn phá quy tắc của chính mình trong tuần này? Trả lời thật, "
     "không ai phán xét.",
     "Pinned comment", "Thấp"),
    ("Psychology / discipline", "Gây tò mò",
     "Có 3 dấu hiệu báo trước rằng bạn sắp revenge trade. Mình từng bỏ qua cả 3 và "
     "mất 1 tài khoản. Bạn đã từng bỏ qua dấu hiệu nào chưa?",
     "Pinned comment", "Thấp"),

    ("Broker / prop firm", "Gợi ý tranh cãi nhẹ",
     "Nhiều người mua prop challenge rồi thất bại không phải vì kém, mà vì đọc "
     "rules chưa kỹ. Bạn đã đọc hết rules trước khi mua chưa?",
     "Pinned comment video prop firm", "Trung bình — cần disclosure nếu có affiliate"),
    ("Broker / prop firm", "Hỏi ngược",
     "Điều kiện nào trong prop challenge làm bạn khó chịu nhất: drawdown tối đa, "
     "ngày giao dịch tối thiểu, hay thời gian giữ lệnh?",
     "Pinned comment", "Thấp"),

    ("Community", "Gợi ý tranh cãi nhẹ",
     "Ý kiến có thể gây tranh cãi: trade một mình là lý do lớn nhất khiến bạn "
     "không tiến bộ. Ai thấy khác, mình muốn nghe.",
     "Pinned comment", "Thấp"),
    ("Community", "Hỏi ngược",
     "Bạn có ai để review lệnh của mình không? Nếu không, đó có thể là phần đang thiếu.",
     "Pinned comment", "Thấp"),

    ("Basics", "Gợi ý tranh cãi nhẹ",
     "Học SMC/ICT trước khi vững basic là thứ tự sai — mình sẵn sàng tranh luận "
     "điều này. Bạn học gì trước?",
     "Pinned comment video basic", "Thấp"),
    ("Basics", "Gây tò mò",
     "Có 4 thứ cơ bản mà nếu bỏ qua, mọi strategy phức tạp đều vô nghĩa. Bạn thử "
     "liệt kê trước xem trùng được mấy cái?",
     "Pinned comment", "Thấp"),
]

REPLY_TEMPLATES = [
    ("Khen chung chung", "\"Cảm ơn bạn!\"", "\"Cảm ơn bạn đã xem. Bạn đang trade "
     "khung thời gian nào? Mình đang tò mò phần lớn khán giả ở đây dùng gì.\""),
    ("Hỏi lại nội dung video", "Trả lời ngay đáp án", "Trả lời + dẫn về mốc thời "
     "gian cụ thể trong video, rồi hỏi ngược 1 câu"),
    ("Chia sẻ đã thua tiền", "Đồng cảm suông", "Đồng cảm + hỏi họ đã xử lý thế nào "
     "+ không hứa sẽ giúp họ có lãi"),
    ("Khen nhưng nghi ngờ", "Xoá / bỏ qua", "Trả lời công khai, thừa nhận giới hạn, "
     "không hạ thấp người khác"),
    ("Comment spam / link lạ", "Trả lời", "Xoá + report. Không trả lời, không bấm link"),
    ("Comment tiêu cực đúng", "Xoá", "Giữ lại + trả lời thừa nhận điểm đúng. Đây là "
     "tín hiệu tốt cho kênh"),
]


def build_policy_rows() -> list[list]:
    return [
        ["Chính kênh đăng comment ghim trên video của mình",
         "ĐƯỢC", "Chủ kênh toàn quyền tương tác trên nội dung của mình"],
        ["Đăng bình luận trên Community tab kêu gọi thảo luận",
         "ĐƯỢC", "Tính năng chính thức của YouTube"],
        ["Trả lời comment khán giả, hỏi ngược để tăng thảo luận",
         "ĐƯỢC", "Tương tác thật, không thao túng"],
        ["Tạo nhiều tài khoản để tự comment khen/seed kênh mình",
         "KHÔNG", "Vi phạm Fake Engagement Policy — có thể bị xoá kênh"],
        ["Mua like / view / comment từ dịch vụ bên ngoài",
         "KHÔNG", "Vi phạm Fake Engagement Policy"],
        ["Rải comment kèm link Telegram/broker trên kênh người khác",
         "KHÔNG", "Spam + Deceptive Practices, dễ bị ban tài khoản"],
        ["Thuê người vào comment có kịch bản dưới danh nghĩa khán giả thật",
         "KHÔNG", "Thao túng tương tác — rủi ro pháp lý và nền tảng"],
        ["Dùng chính kênh của mình đăng comment trên kênh đối thủ có giá trị thật",
         "THẬN TRỌNG", "Chỉ khi đóng góp thật, không link, không quảng cáo"],
        ["Kêu gọi khán giả comment 1 từ khoá (RISK/PLAN/ENTRY)",
         "ĐƯỢC", "CTA thông thường, miễn không hứa lợi nhuận"],
        ["Cam kết lợi nhuận hoặc show trade giả trong comment",
         "KHÔNG", "Vi phạm guardrail dự án + chính sách quảng cáo tài chính"],
    ]


# ── DOCX ────────────────────────────────────────────────────────────────────

def build_docx(pain_all, backlog, angles):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(4)

    t = doc.add_heading("ACTION PLAN — ĐỘI EDIT & BỘ COMMENT TƯƠNG TÁC", level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Kênh @azzammastertradinggold  •  Ngách XAUUSD / Forex")
    r.bold = True
    r.font.size = Pt(12)
    m = doc.add_paragraph()
    m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    m.add_run(f"Ngày lập: {date.today().isoformat()}").italic = True
    doc.add_paragraph()

    # KPI
    doc.add_heading("1. KPI NGÀY", level=1)
    kpi_tbl = doc.add_table(rows=4, cols=2)
    kpi_tbl.style = "Light Grid Accent 1"
    pairs = [
        ("Short video / ngày", f"{KPI['shorts_per_day']} video (mỗi video 15-35 giây)"),
        ("Long video / ngày", f"{KPI['long_per_day']} video (~{KPI['long_minutes']} phút)"),
        ("Tổng sản lượng / tuần", f"{KPI['shorts_per_day']*7} Short + {KPI['long_per_day']*7} Long"),
        ("Nguồn nguyên liệu", f"{len(pain_all):,} pain point unique + {len(backlog)} content backlog"),
    ]
    for i, (k, v) in enumerate(pairs):
        kpi_tbl.cell(i, 0).text = k
        kpi_tbl.cell(i, 1).text = v
        for p in kpi_tbl.cell(i, 0).paragraphs:
            for run in p.runs:
                run.bold = True
    doc.add_paragraph()

    # Warning about policy
    doc.add_heading("2. RANH GIỚI CHÍNH SÁCH YOUTUBE — ĐỌC TRƯỚC KHI LÀM", level=1)
    warn = doc.add_paragraph()
    wr = warn.add_run(
        "Bộ comment trong tài liệu này KHÔNG phải để tài khoản ảo đi rải. "
        "Nó dùng cho 2 việc hợp lệ: (a) chính kênh đăng comment ghim trên video "
        "của mình, (b) khán giả thật dùng để thảo luận. Mọi hình thức tạo tài "
        "khoản giả, mua engagement, hay rải link sang kênh khác đều vi phạm "
        "Fake Engagement Policy và có thể dẫn tới xoá kênh."
    )
    wr.bold = True
    wr.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
    doc.add_paragraph()

    for title, rows in [("2.1 Được phép / Không được phép", build_policy_rows())]:
        doc.add_heading(title, level=2)
        tb = doc.add_table(rows=len(rows) + 1, cols=3)
        tb.style = "Light Grid Accent 1"
        for j, h in enumerate(["Hành vi", "Kết luận", "Căn cứ"]):
            c = tb.cell(0, j)
            c.text = h
            for p in c.paragraphs:
                for run in p.runs:
                    run.bold = True
        for i, row in enumerate(rows, 1):
            for j, val in enumerate(row):
                tb.cell(i, j).text = val
        doc.add_paragraph()

    # Workflow
    doc.add_heading("3. WORKFLOW SẢN XUẤT", level=1)
    wf = build_workflow_rows()
    tb = doc.add_table(rows=len(wf) + 1, cols=4)
    tb.style = "Light Grid Accent 1"
    for j, h in enumerate(["Bước", "Ai làm", "Việc cụ thể", "Output"]):
        c = tb.cell(0, j)
        c.text = h
        for p in c.paragraphs:
            for run in p.runs:
                run.bold = True
    for i, row in enumerate(wf, 1):
        for j, val in enumerate(row):
            tb.cell(i, j).text = val
    doc.add_paragraph()

    # KPI detail
    doc.add_heading("4. PHÂN BỔ KPI THEO TỪNG VIDEO", level=1)
    kr = build_kpi_rows(pain_all, backlog)
    tb = doc.add_table(rows=len(kr) + 1, cols=6)
    tb.style = "Light Grid Accent 1"
    for j, h in enumerate(["Slot", "Loại nội dung", "Thời lượng", "Ràng buộc",
                           "KPI chính", "Nguồn dữ liệu"]):
        c = tb.cell(0, j)
        c.text = h
        for p in c.paragraphs:
            for run in p.runs:
                run.bold = True
    for i, row in enumerate(kr, 1):
        for j, val in enumerate(row):
            tb.cell(i, j).text = val
    doc.add_paragraph()

    # Data sources
    doc.add_heading("5. NGUỒN DỮ LIỆU PAIN POINT — LẤY Ở ĐÂU, CỘT NÀO", level=1)
    sr = build_source_rows()
    tb = doc.add_table(rows=len(sr) + 1, cols=4)
    tb.style = "Light Grid Accent 1"
    for j, h in enumerate(["Loại dữ liệu", "Sheet trong Excel", "Cột cần đọc", "Dùng để"]):
        c = tb.cell(0, j)
        c.text = h
        for p in c.paragraphs:
            for run in p.runs:
                run.bold = True
    for i, row in enumerate(sr, 1):
        for j, val in enumerate(row):
            tb.cell(i, j).text = val
    doc.add_paragraph()

    doc.add_heading("5.1 Ví dụ lấy nguyên liệu (làm đúng theo bước này)", level=2)
    doc.add_paragraph(
        "1. Mở file Excel AZZAM_PHAN_TICH_DOI_THU.xlsx → sheet 09_PAINPOINT_ALL."
    )
    doc.add_paragraph(
        "2. Bấm filter ở cột Score, chọn Top 10 (score cao nhất)."
    )
    doc.add_paragraph(
        "3. Đọc cột Categories để chọn 3 comment thuộc 3 nhóm pain KHÁC nhau — "
        "không chọn 3 comment cùng nhóm, vì 3 Short trong ngày phải phủ 3 nỗi đau khác nhau."
    )
    doc.add_paragraph(
        "4. Copy nguyên văn comment ở cột Nội dung comment. Đây là câu khán giả "
        "thật viết — dùng làm hook, không cần sáng tạo lại."
    )
    doc.add_paragraph(
        "5. Ghi lại Comment ID + URL nguồn vào file theo dõi của editor để audit."
    )
    doc.add_paragraph(
        "6. Với Long video: mở sheet 04_SALES_ANGLES, chọn 1 pillar, đọc Big promise "
        "+ Proof hook để viết outline 8 phút."
    )
    doc.add_paragraph()

    # Seed comments
    doc.add_heading("6. BỘ COMMENT TƯƠNG TÁC", level=1)
    doc.add_paragraph(
        "Nguyên tắc: câu hỏi phải trả lời được bằng 1-2 câu, có quan điểm rõ để "
        "người khác muốn phản hồi, và không hứa hẹn kết quả trading. Đăng ở "
        "pinned comment hoặc Community tab — không dùng tài khoản khác."
    )
    doc.add_paragraph()
    current = None
    for grp, kind, text, where, risk in SEED_COMMENTS:
        if grp != current:
            doc.add_heading(f"Nhóm: {grp}", level=2)
            current = grp
        p = doc.add_paragraph(style="List Bullet")
        r1 = p.add_run(f"[{kind}] ")
        r1.bold = True
        p.add_run(text)
        meta = doc.add_paragraph()
        meta.paragraph_format.left_indent = Pt(24)
        mr = meta.add_run(f"→ Dùng ở: {where}  |  Rủi ro chính sách: {risk}")
        mr.italic = True
        mr.font.size = Pt(9)
        mr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph()
    doc.add_heading("6.1 Mẫu trả lời comment khán giả", level=2)
    tb = doc.add_table(rows=len(REPLY_TEMPLATES) + 1, cols=3)
    tb.style = "Light Grid Accent 1"
    for j, h in enumerate(["Loại comment", "Cách làm SAI", "Cách làm ĐÚNG"]):
        c = tb.cell(0, j)
        c.text = h
        for p in c.paragraphs:
            for run in p.runs:
                run.bold = True
    for i, row in enumerate(REPLY_TEMPLATES, 1):
        for j, val in enumerate(row):
            tb.cell(i, j).text = val
    doc.add_paragraph()

    # Measurement
    doc.add_heading("7. ĐO LƯỜNG HẰNG NGÀY", level=1)
    for txt in [
        f"Sản lượng: đủ {KPI['shorts_per_day']} Short + {KPI['long_per_day']} Long/ngày chưa?",
        "Chất lượng hook: tỷ lệ giữ chân 3 giây đầu (mục tiêu >60%).",
        "Long video: thời lượng xem trung bình (mục tiêu >40%).",
        "Tương tác: số comment trên mỗi video, số trả lời của kênh.",
        "Lead: số người vào Telegram từ comment keyword.",
        "Tuân thủ: có video nào cam kết lợi nhuận / trade giả không? (phải bằng 0)",
    ]:
        doc.add_paragraph(txt, style="List Bullet")
    doc.add_paragraph()

    doc.add_heading("8. VIỆC CẦN DUYỆT TRƯỚC KHI CHẠY", level=1)
    for txt in [
        "Xác nhận đội edit gồm mấy người, ai chịu trách nhiệm slot nào.",
        "Xác nhận ngôn ngữ: tiếng Việt hay tiếng Anh (ảnh hưởng toàn bộ hook + title).",
        "Xác nhận kênh Telegram đích để gắn CTA.",
        "Xác nhận có dùng affiliate broker/prop firm chưa (để thêm disclosure).",
        "Duyệt bộ comment tương tác trước khi đăng lần đầu.",
    ]:
        doc.add_paragraph(txt, style="List Number")

    path = OUT / "AZZAM_ACTION_PLAN_EDITOR.docx"
    doc.save(path)
    return path


# ── XLSX ────────────────────────────────────────────────────────────────────

def build_xlsx(pain_all, backlog, angles):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    HF = PatternFill("solid", fgColor="1F3864")
    HFONT = Font(bold=True, color="FFFFFF", size=10)
    TFONT = Font(bold=True, size=14, color="1F3864")
    NFONT = Font(italic=True, size=9, color="555555")
    THIN = Side(style="thin", color="BFBFBF")
    BD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    wb = Workbook()

    def sheet(name, title, note, headers, rows, widths=None, numbers=None, wrap=None):
        ws = wb.create_sheet(name) if wb.sheetnames != ["Sheet"] or ws_used[0] else wb.active
        ws_used[0] = True
        ws.title = name
        ws["A1"] = title
        ws["A1"].font = TFONT
        ws["A2"] = note
        ws["A2"].font = NFONT
        hr = 4
        for j, h in enumerate(headers, 1):
            c = ws.cell(row=hr, column=j, value=h)
            c.fill = HF
            c.font = HFONT
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BD
        for i, row in enumerate(rows, hr + 1):
            for j, v in enumerate(row, 1):
                c = ws.cell(row=i, column=j, value=v)
                c.border = BD
                c.font = Font(size=10)
                c.alignment = Alignment(vertical="top",
                                        wrap_text=(headers[j - 1] in (wrap or set())))
                if numbers and headers[j - 1] in numbers and isinstance(v, (int, float)):
                    c.number_format = "#,##0"
        for j, h in enumerate(headers, 1):
            L = get_column_letter(j)
            if widths and h in widths:
                ws.column_dimensions[L].width = widths[h]
            else:
                longest = len(str(h))
                for row in rows[:200]:
                    if j - 1 < len(row):
                        longest = max(longest, len(str(row[j - 1])))
                ws.column_dimensions[L].width = min(max(longest + 2, 10), 70)
        if rows:
            ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(headers))}{hr + len(rows)}"
        ws.freeze_panes = ws.cell(row=hr + 1, column=1)
        return ws

    ws_used = [False]

    # 01 KPI
    ws = wb.active
    ws.title = "01_KPI_NGAY"
    ws["A1"] = "KPI NGÀY — ĐỘI EDIT"
    ws["A1"].font = TFONT
    ws["A2"] = (f"{KPI['shorts_per_day']} Short + {KPI['long_per_day']} Long "
                f"(~{KPI['long_minutes']} phút) mỗi ngày")
    ws["A2"].font = NFONT
    headers = ["Slot", "Loại nội dung", "Thời lượng", "Ràng buộc", "KPI chính", "Nguồn dữ liệu"]
    rows = build_kpi_rows(pain_all, backlog)
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.fill = HF; c.font = HFONT; c.border = BD
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, row in enumerate(rows, 5):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v); c.border = BD; c.font = Font(size=10)
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for j, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(j)].width = [10, 24, 16, 26, 24, 40][j - 1]
    ws.freeze_panes = "A5"
    ws_used[0] = True

    # 02 WORKFLOW
    sheet("02_WORKFLOW", "WORKFLOW SẢN XUẤT",
          "Thứ tự bắt buộc — không nhảy bước",
          ["Bước", "Ai làm", "Việc cụ thể", "Output"],
          build_workflow_rows(),
          widths={"Bước": 22, "Ai làm": 16, "Việc cụ thể": 55, "Output": 38},
          wrap={"Việc cụ thể", "Output"})

    # 03 DATA SOURCE
    sheet("03_NGUON_DU_LIEU", "NGUỒN DỮ LIỆU PAIN POINT",
          "Đội edit tự lấy nguyên liệu từ các sheet này trong AZZAM_PHAN_TICH_DOI_THU.xlsx",
          ["Loại dữ liệu", "Sheet trong Excel", "Cột cần đọc", "Dùng để"],
          build_source_rows(),
          widths={"Loại dữ liệu": 26, "Sheet trong Excel": 24,
                  "Cột cần đọc": 52, "Dùng để": 34},
          wrap={"Cột cần đọc", "Dùng để"})

    # 04 PAIN POINT POOL
    pool = sorted(pain_all, key=lambda x: -to_float(x.get("score")))[:200]
    rows = []
    for i, p in enumerate(pool, 1):
        rows.append([
            i, to_float(p.get("score")), p.get("categories", ""),
            p.get("engagement_likes", 0),
            unescape(p.get("comment_text", ""))[:400],
            p.get("content_url", ""), p.get("comment_id", ""),
        ])
    sheet("04_PAINPOOL_200", "200 PAIN POINT ƯU TIÊN CAO NHẤT",
          f"Lấy từ {len(pain_all):,} pain point unique. Dùng cột Categories để chọn "
          "3 nhóm khác nhau cho 3 Short.",
          ["#", "Score", "Categories", "Likes", "Nội dung comment (nguyên liệu hook)",
           "URL nguồn", "Comment ID"],
          rows,
          widths={"#": 6, "Score": 9, "Categories": 34, "Likes": 8,
                  "Nội dung comment (nguyên liệu hook)": 90,
                  "URL nguồn": 42, "Comment ID": 26},
          numbers={"Score", "Likes"}, wrap={"Nội dung comment (nguyên liệu hook)"})

    # 05 SEED COMMENTS
    rows = [[g, k, t, w, r] for g, k, t, w, r in SEED_COMMENTS]
    sheet("05_COMMENT_TUONG_TAC", "BỘ COMMENT TƯƠNG TÁC",
          "CHỈ dùng cho pinned comment trên video của kênh mình hoặc Community tab. "
          "KHÔNG dùng tài khoản khác để rải.",
          ["Nhóm", "Loại", "Câu comment", "Dùng ở đâu", "Rủi ro chính sách"],
          rows,
          widths={"Nhóm": 26, "Loại": 22, "Câu comment": 85,
                  "Dùng ở đâu": 30, "Rủi ro chính sách": 38},
          wrap={"Câu comment", "Rủi ro chính sách"})

    # 06 REPLY TEMPLATES
    rows = [[a, b, c] for a, b, c in REPLY_TEMPLATES]
    sheet("06_MAU_TRA_LOI", "MẪU TRẢ LỜI COMMENT KHÁN GIẢ",
          "Trả lời đúng cách giữ được thảo luận và không tạo rủi ro.",
          ["Loại comment", "Cách làm SAI", "Cách làm ĐÚNG"],
          rows,
          widths={"Loại comment": 30, "Cách làm SAI": 42, "Cách làm ĐÚNG": 80},
          wrap={"Cách làm SAI", "Cách làm ĐÚNG"})

    # 07 POLICY
    sheet("07_CHINH_SACH", "RANH GIỚI CHÍNH SÁCH YOUTUBE",
          "Vi phạm Fake Engagement Policy có thể dẫn tới xoá kênh — đọc kỹ cột Kết luận.",
          ["Hành vi", "Kết luận", "Căn cứ"],
          build_policy_rows(),
          widths={"Hành vi": 62, "Kết luận": 14, "Căn cứ": 58},
          wrap={"Hành vi", "Căn cứ"})

    # 08 LONG VIDEO OUTLINE
    rows = []
    for a in angles:
        rows.append([
            a["id"], a["pain_cluster"], a.get("evidence_count", 0),
            a.get("big_promise", ""), a.get("proof_hook", ""), a.get("cta", ""),
            " | ".join(a.get("content_formats", [])),
        ])
    sheet("08_LONG_OUTLINE", "OUTLINE LONG VIDEO 8 PHÚT",
          "Chọn 1 pillar mỗi ngày. Cấu trúc: vấn đề → nguyên nhân → giải pháp → "
          "checklist → CTA.",
          ["Pillar", "Pain cluster", "Evidence", "Big promise", "Proof hook",
           "CTA", "Format đề xuất"],
          rows,
          widths={"Pillar": 8, "Pain cluster": 34, "Evidence": 10,
                  "Big promise": 50, "Proof hook": 50, "CTA": 30,
                  "Format đề xuất": 45},
          numbers={"Evidence"}, wrap={"Big promise", "Proof hook", "Format đề xuất"})

    # 09 TRACKING
    sheet("09_THEO_DOI", "BẢNG THEO DÕI SẢN XUẤT (điền mỗi ngày)",
          "Điền trực tiếp trên sheet này. Mỗi dòng = 1 video.",
          ["Ngày", "Slot", "Loại", "Tiêu đề", "Pain point gốc (Comment ID)",
           "URL nguồn", "Hook đã viết", "Trạng thái", "Người làm", "Ghi chú"],
          [["", "", "", "", "", "", "", "chưa làm", "", ""] for _ in range(30)],
          widths={"Ngày": 12, "Slot": 9, "Loại": 9, "Tiêu đề": 40,
                  "Pain point gốc (Comment ID)": 26, "URL nguồn": 38,
                  "Hook đã viết": 50, "Trạng thái": 13, "Người làm": 14,
                  "Ghi chú": 30},
          wrap={"Tiêu đề", "Hook đã viết", "Ghi chú"})

    path = OUT / "AZZAM_ACTION_PLAN_EDITOR.xlsx"
    wb.save(path)
    return path


def main() -> int:
    pp = dedupe(load_csv("outputs/competitor_analysis/painpoint_candidates.csv")
                + load_csv("outputs/competitor_longform/painpoint_candidates.csv"))
    backlog = load_csv("outputs/strategy/content_backlog.csv")
    angles = load_json("outputs/strategy/sales_angles.json") or []

    print(f"Pain points unique: {len(pp)}")
    print(f"Backlog: {len(backlog)} | Angles: {len(angles)}")
    print(f"Seed comments: {len(SEED_COMMENTS)} | Reply templates: {len(REPLY_TEMPLATES)}")

    d = build_docx(pp, backlog, angles)
    x = build_xlsx(pp, backlog, angles)

    for p in (d, x):
        print(f"Saved: {p.resolve()}")
        print(f"  size: {p.stat().st_size:,} bytes")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
